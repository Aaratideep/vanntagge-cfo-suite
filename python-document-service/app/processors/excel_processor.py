import openpyxl
from app.processors.base_processor import BaseProcessor
from app.models.processing_result import ProcessingResult, DocumentMetadata

class ExcelProcessor(BaseProcessor):
    def process(self, file_path: str, document_id: str, document_type: str, file_name: str, file_size: int) -> ProcessingResult:
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            
            total_rows = 0
            tables = []
            
            for sheet_name in sheet_names:
                sheet = wb[sheet_name]
                # In read_only mode, we can iterate rows
                row_count = 0
                headers = []
                sample_data = []
                
                for row in sheet.iter_rows(values_only=True):
                    row_count += 1
                    if row_count == 1:
                        headers = [str(cell) if cell is not None else "" for cell in row]
                    elif row_count <= 5:
                        sample_data.append([str(cell) if cell is not None else "" for cell in row])
                
                total_rows += row_count
                
                tables.append({
                    "sheetName": sheet_name,
                    "rows": row_count,
                    "columns": len(headers) if headers else 0,
                    "headers": headers,
                    "sample": sample_data
                })
                
            wb.close()
            
            return ProcessingResult(
                documentId=document_id,
                documentType=document_type,
                processor="excel",
                status="COMPLETED",
                metadata=DocumentMetadata(
                    fileName=file_name,
                    fileSizeBytes=file_size,
                    sheetCount=len(sheet_names),
                    rowCount=total_rows
                ),
                tables=tables
            )
        except Exception as e:
            return ProcessingResult(
                documentId=document_id,
                documentType=document_type,
                processor="excel",
                status="FAILED",
                metadata=DocumentMetadata(
                    fileName=file_name,
                    fileSizeBytes=file_size
                ),
                errorMessage=str(e)
            )
