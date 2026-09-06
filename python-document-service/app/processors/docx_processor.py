import docx
from app.processors.base_processor import BaseProcessor
from app.models.processing_result import ProcessingResult, DocumentMetadata

class DocxProcessor(BaseProcessor):
    def process(self, file_path: str, document_id: str, document_type: str, file_name: str, file_size: int) -> ProcessingResult:
        try:
            doc = docx.Document(file_path)
            
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            extracted_text = "\n".join(paragraphs)
            
            tables = []
            for i, table in enumerate(doc.tables):
                table_data = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_data.append(row_data)
                
                if table_data:
                    headers = table_data[0]
                    sample = table_data[1:6] if len(table_data) > 1 else []
                    tables.append({
                        "sheetName": f"Table {i+1}",
                        "rows": len(table_data),
                        "columns": len(headers),
                        "headers": headers,
                        "sample": sample
                    })
            
            return ProcessingResult(
                documentId=document_id,
                documentType=document_type,
                processor="docx",
                status="COMPLETED",
                metadata=DocumentMetadata(
                    fileName=file_name,
                    fileSizeBytes=file_size,
                    pageCount=None # DOCX doesn't easily expose page count without rendering
                ),
                extractedText=extracted_text[:50000],
                tables=tables
            )
        except Exception as e:
            return ProcessingResult(
                documentId=document_id,
                documentType=document_type,
                processor="docx",
                status="FAILED",
                metadata=DocumentMetadata(
                    fileName=file_name,
                    fileSizeBytes=file_size
                ),
                errorMessage=str(e)
            )
